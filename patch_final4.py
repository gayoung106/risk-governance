import os, shutil
from docx import Document

def sub_runs(para, old, new):
    full = para.text
    if old not in full:
        return False
    changed = False
    for run in para.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            changed = True
    if changed:
        return True
    new_full = full.replace(old, new)
    if para.runs:
        para.runs[0].text = new_full
        for r in para.runs[1:]:
            r.text = ''
    return True

def patch_cell(cell, old, new):
    for p in cell.paragraphs:
        sub_runs(p, old, new)

doc = Document('result/final_text_submission_ready.docx')
paras = doc.paragraphs
P = []  # patches log

# ═══════════════════════════════════════════════════════
# [수정 1] coefficient superiority / dependence 제거
# ═══════════════════════════════════════════════════════

# Para [12]: 핵심 기반 → 유의하게 연관
if sub_runs(paras[12],
    "관계적 신뢰가 정당성 인식의 핵심 기반이 되는가",
    "관계적 신뢰가 정당성 인식과 유의하게 연관되는가"):
    P.append('[12] 핵심 기반 → 유의하게 연관')

# Para [36]: 더 의존하는 경향 → 유의하게 연관될 것
if sub_runs(paras[36],
    "형식적 안전관리 인식보다 관계적으로 구축된 관리신뢰에 더 의존하는 경향을 이론적으로 예측한다",
    "관계적으로 구축된 관리신뢰와 안전관리 인식이 모두 정책수용성과 유의하게 연관될 것을 이론적으로 예측한다"):
    P.append('[36] 더 의존하는 경향 → 유의하게 연관')

# Para [106]: 강한 정적 → 유의한 정적
if sub_runs(paras[106],
    "정책수용성과 강한 정적 연관을 나타냈다",
    "정책수용성과 유의한 정적 연관을 나타냈다"):
    P.append('[106] 강한 정적 → 유의한 정적')

# Para [110]: 강한 정(+) → 유의한 정(+)
if sub_runs(paras[110],
    "정책수용성과 강한 정(+)의 연관성이 확인되며",
    "정책수용성과 유의한 정(+)의 연관성이 확인되며"):
    P.append('[110] 강한 정(+) → 유의한 정(+)')

# Para [119]: 강한 정적 연관성 → 유의한 정적 연관성
if sub_runs(paras[119],
    "제도적 신뢰와 강한 정적 연관성을 나타냈다(β=+0.778",
    "제도적 신뢰와 유의한 정적 연관성을 나타냈다(β=+0.778"):
    P.append('[119] 강한 정적 → 유의한 정적 (Path B)')

# Para [131]: 강한 정적 + 더욱 결정적 → 유의한 + 더욱 뚜렷하게
if sub_runs(paras[131],
    "관리신뢰는 정책수용성과 강한 정적 연관성을 나타내지만, 위험 상황에서 관리신뢰가 정당성 자원으로서 더욱 결정적인 역할을 수행함을 구체적으로 보여준다",
    "관리신뢰는 정책수용성과 유의한 정적 연관성을 나타내지만, 위험 상황에서 이 연관성이 더욱 뚜렷하게 나타남을 구체적으로 보여준다"):
    P.append('[131] 강한 정적+결정적 역할 → 유의한+뚜렷하게')

# Para [147]: 강한 정적 → 유의한 정적
if sub_runs(paras[147],
    "확산적 제도적 신뢰와 강한 정적 연관성을 나타낸다는 것이다",
    "확산적 제도적 신뢰와 유의한 정적 연관성을 나타낸다는 것이다"):
    P.append('[147] 강한 정적 → 유의한 정적 (Discussion)')

# Para [151]: 핵심 자원 → 관련 자원
if sub_runs(paras[151],
    "관계적 신뢰가 거버넌스 정당성의 핵심 자원이라는 명제",
    "관계적 신뢰가 거버넌스 정당성과 유의하게 연관된다는 명제"):
    P.append('[151] 핵심 자원 → 유의하게 연관')

# Table T6R1C2: 강한 정적 → 유의한 정적 (simple slopes table)
cell = doc.tables[6].cell(1, 2)
if '강한 정적 연관' in cell.text:
    patch_cell(cell, '강한 정적 연관', '유의한 정적 연관')
    P.append('T6R1C2 강한 정적 → 유의한 정적')

# ═══════════════════════════════════════════════════════
# [수정 2] legitimacy production / authorization 완화
# ═══════════════════════════════════════════════════════

# Para [4]: 정당성 생산 조건 → 정당성 인식 조건
if sub_runs(paras[4],
    "사후규제 체계의 정당성 생산 조건에 관한 후속 비교·종단 연구를 위한 경험적 출발점을 제시한다",
    "사후규제 체계의 정당성 인식 조건에 관한 후속 비교·종단 연구를 위한 경험적 출발점을 제시한다"):
    P.append('[4] 정당성 생산 조건 → 정당성 인식 조건 (abstract)')

# Para [22]: 거버넌스 정당성 생산 (x2)
if sub_runs(paras[22],
    "실질적 거버넌스 정당성 생산보다 책임 회피",
    "실질적 거버넌스 정당성 인식 형성보다 책임 회피"):
    P.append('[22] 정당성 생산 → 정당성 인식 형성 (Hood)')
if sub_runs(paras[22],
    "확산적 정당성 생산 역량을 약화시키는",
    "확산적 정당성 인식 형성 역량을 약화시키는"):
    P.append('[22] 확산적 정당성 생산 → 정당성 인식 형성')

# Para [118]: 직접 정당성 수권 → 직접 정당성 인식
if sub_runs(paras[118],
    "경로 1(직접 정당성 수권)",
    "경로 1(직접 정당성 인식)"):
    P.append('[118] 정당성 수권 → 정당성 인식')

# Para [145]: 거버넌스 정당성 생산의 구조적 조건
if sub_runs(paras[145],
    "거버넌스 정당성 생산의 구조적 조건에 관한 검증 가능한 명제들을 도출하는",
    "거버넌스 정당성 인식의 구조적 조건에 관한 탐색적 명제들을 도출하는"):
    P.append('[145] 정당성 생산의 구조적 조건 → 정당성 인식의 구조적 조건')

# Para [145]: 정당성 생산 기제
if sub_runs(paras[145],
    "정당성 생산 기제로서의 효율을 체계적으로 상실한다",
    "정당성 인식 형성 기제로서의 효율을 상실하는 경향이 있다"):
    P.append('[145] 정당성 생산 기제 → 정당성 인식 형성 기제')

# Para [145]: "형식적 안전관리 신호의 부적 연관성—은 이 이론적 예측의 경험적 표현이다"
# → positive finding 반영, causal claim 제거
if sub_runs(paras[145],
    "본 연구의 발견—형식적 안전관리 신호의 부적 연관성—은 이 이론적 예측의 경험적 표현이다.",
    "본 연구의 탐색적 발견—관리신뢰의 유의한 정적 연관성 패턴—은 이 이론적 시각과의 탐색적 접점을 제공한다."):
    P.append('[145] 부적 연관성 경험적 표현 → 탐색적 접점 (방향 수정)')

# Para [146]: 정당성 이전 → 정당성 인식 연관
if sub_runs(paras[146],
    "신뢰 자원으로의 정당성 이전",
    "신뢰 자원과의 정당성 인식 연관"):
    P.append('[146] 정당성 이전 → 정당성 인식 연관')

# Para [156]: 정당성 수권 효과 → 정책수용성 간 연관
if sub_runs(paras[156],
    "관리신뢰의 정당성 수권 효과가 증폭됨을 보여준다",
    "관리신뢰와 정책수용성 간 유의한 연관이 위험인식 수준에 따라 증폭됨을 보여준다"):
    P.append('[156] 정당성 수권 효과 → 유의한 연관 증폭')

# Para [168]: 정당성 생산 조건 → 정당성 인식 조건
if sub_runs(paras[168],
    "사후규제 거버넌스의 정당성 생산 조건에 관한 보다 확고한 이론적 기반",
    "사후규제 거버넌스의 정당성 인식 조건에 관한 탐색적 이론적 기반"):
    P.append('[168] 정당성 생산 조건 → 정당성 인식 조건 (conclusion)')

# ═══════════════════════════════════════════════════════
# [수정 3] safety variable 방향 충돌 최종 제거
# (para[145] 이미 수정. 나머지 잔존 부적 검사)
# ═══════════════════════════════════════════════════════
# Para [111] H2: risk→consent 부적은 정당(유지)
# Para [166] H2: risk→consent 부적은 정당(유지)
# No action needed for H2 instances

# ═══════════════════════════════════════════════════════
# [수정 4] abstract / conclusion claim intensity 완화
# ═══════════════════════════════════════════════════════

# Para [145]: 핵심 이론적 기여 → 탐색적 이론적 함의
if sub_runs(paras[145],
    "본 연구의 핵심 이론적 기여는",
    "본 연구의 탐색적 이론적 함의는"):
    P.append('[145] 핵심 이론적 기여 → 탐색적 이론적 함의')

# Para [4]: 거버넌스 정당성 인식 연구에 탐색적 이론적 기반 제공 - check if already OK
# Already says "탐색적 이론적 기반" so no change needed

# ═══════════════════════════════════════════════════════
# 저장
# ═══════════════════════════════════════════════════════
doc.save('result/final_text_patched4.tmp')
shutil.move('result/final_text_patched4.tmp', 'result/final_text_patched4.docx')

print(f'Patches applied: {len(P)}')
for p in P:
    print(f'  ✓ {p}')

# ═══════════════════════════════════════════════════════
# [수정 5] consistency audit
# ═══════════════════════════════════════════════════════
print()
print('=== CONSISTENCY AUDIT ===')
doc2 = Document('result/final_text_patched4.docx')

audit_bad = [
    # Superiority
    '강한 정적 연관', '강한 정(+)', '더욱 결정적', '핵심 자원이라는', '핵심 기반이 되는',
    '더 의존하는', 'dominant', 'stronger', 'superiority',
    # Legitimacy production
    '정당성 생산', '정당성 수권', '정당성 이전',
    # Paradox
    '역설적', '역설 패턴', 'paradox',
    # Mediation causal
    '부분적으로 매개', '매개 역할', '매개효과',
    # Wrong direction
    '안전관리 인식이 정책수용성(β=−', '안전관리 인식(β=−',
    '부적 연관성—', '부적 연관성은 이',
    # Variable name
    'safety_perception',
    # Coefficient comparison
    '비대칭 패턴', '에 비해 크게',
]
all_ok = True
for term in audit_bad:
    locs = []
    for i, p in enumerate(doc2.paragraphs):
        if term in p.text:
            locs.append(f'para[{i}]')
    for ti, t in enumerate(doc2.tables):
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                if term in cell.text:
                    locs.append(f'T{ti}R{ri}C{ci}')
    if locs:
        print(f'  WARNING: "{term}" → {locs}')
        all_ok = False
    else:
        print(f'  OK: "{term}"')

# Positive direction verification
pos_checks = [(113,'β=+0.191'),(119,'β=+0.778'),(134,'coef=+0.636')]
print()
for idx, term in pos_checks:
    status = '✓' if term in doc2.paragraphs[idx].text else '✗ MISSING'
    print(f'  {status} {term} in para[{idx}]')

print()
if all_ok:
    print('ALL CLEAR — final_text_patched4.docx SUBMISSION-READY')
else:
    print('일부 항목 남음 — WARNING 확인 요망')
print(f'File: {os.path.getsize("result/final_text_patched4.docx"):,} bytes')
